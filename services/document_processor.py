import os
import aiofiles
from pathlib import Path
from typing import List, Dict, Any
import PyPDF2
import docx
from openpyxl import load_workbook
from config import settings


class DocumentProcessor:
    
    @staticmethod
    async def extract_text(file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        
        if ext == '.pdf':
            return await DocumentProcessor._extract_pdf(file_path)
        elif ext in ['.docx', '.doc']:
            return await DocumentProcessor._extract_docx(file_path)
        elif ext in ['.txt', '.md']:
            return await DocumentProcessor._extract_text_file(file_path)
        elif ext in ['.xlsx', '.xls']:
            return await DocumentProcessor._extract_excel(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
    
    @staticmethod
    async def _extract_pdf(file_path: str) -> str:
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        return text
    
    @staticmethod
    async def _extract_docx(file_path: str) -> str:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    
    @staticmethod
    async def _extract_text_file(file_path: str) -> str:
        async with aiofiles.open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return await f.read()
    
    @staticmethod
    async def _extract_excel(file_path: str) -> str:
        wb = load_workbook(file_path, read_only=True)
        text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text.append(' | '.join([str(cell) if cell else '' for cell in row]))
        return "\n".join(text)
    
    @staticmethod
    def chunk_text(text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        if chunk_size is None:
            chunk_size = settings.chunk_size
        if overlap is None:
            overlap = settings.chunk_overlap
        
        words = text.split()
        chunks = []
        
        for i in range(0, len(words), chunk_size - overlap):
            chunk = ' '.join(words[i:i + chunk_size])
            if chunk:
                chunks.append(chunk)
        
        return chunks
    
    @staticmethod
    def get_file_metadata(file_path: str) -> Dict[str, Any]:
        stat = os.stat(file_path)
        return {
            'size': stat.st_size,
            'created': stat.st_ctime,
            'modified': stat.st_mtime,
            'extension': Path(file_path).suffix.lower()
        }
